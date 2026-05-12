"""Export service: converts DuckDB table data to CSV, Excel, and JSON bytes.

All public functions accept a db_path and table_name, validate their
existence, query the full table, and return the serialised result as bytes.
They raise ValueError if the database file or named table does not exist.
"""

from __future__ import annotations

import csv
import io
import json
import os
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

import duckdb
import openpyxl
from openpyxl.styles import Font


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------


def _assert_db_and_table(db_path: str, table_name: str) -> None:
    """Raise ValueError if db_path or table_name do not exist."""
    if not os.path.isfile(db_path):
        raise ValueError(f"Database file not found: {db_path}")

    conn = duckdb.connect(db_path, read_only=True)
    try:
        existing: set[str] = {
            row[0]
            for row in conn.execute(
                "SELECT table_name FROM information_schema.tables "
                "WHERE table_schema = 'main'"
            ).fetchall()
        }
    finally:
        conn.close()

    if table_name not in existing:
        raise ValueError(f"Table not found: {table_name}")


def _fetch_table(
    db_path: str,
    table_name: str,
) -> tuple[list[str], list[tuple[Any, ...]]]:
    """Return (column_names, rows) for the given table."""
    conn = duckdb.connect(db_path, read_only=True)
    try:
        result = conn.execute(f'SELECT * FROM "{table_name}"')
        columns: list[str] = [desc[0] for desc in result.description]
        rows: list[tuple[Any, ...]] = result.fetchall()
        return columns, rows
    finally:
        conn.close()


def _to_json_safe(value: Any) -> Any:
    """Convert non-JSON-native Python values to JSON-serialisable equivalents."""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def export_to_csv(db_path: str, table_name: str) -> bytes:
    """Export *table_name* from *db_path* as UTF-8 CSV bytes.

    The first row is a header containing column names.  The csv stdlib
    module is used for serialisation (no DuckDB COPY TO on disk).

    Parameters
    ----------
    db_path:
        Absolute path to the workspace ``.db`` file.
    table_name:
        Name of the table to export.

    Returns
    -------
    bytes
        UTF-8-encoded CSV content including the header row.

    Raises
    ------
    ValueError
        If *db_path* does not exist or *table_name* is not present.
    """
    _assert_db_and_table(db_path, table_name)
    columns, rows = _fetch_table(db_path, table_name)

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(columns)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def export_to_excel(db_path: str, table_name: str) -> bytes:
    """Export *table_name* from *db_path* as .xlsx bytes via openpyxl.

    The worksheet name is *table_name* truncated to 31 characters (the
    Excel limit).  The header row is rendered in bold.

    Parameters
    ----------
    db_path:
        Absolute path to the workspace ``.db`` file.
    table_name:
        Name of the table to export.

    Returns
    -------
    bytes
        Raw bytes of a valid .xlsx workbook.

    Raises
    ------
    ValueError
        If *db_path* does not exist or *table_name* is not present.
    """
    _assert_db_and_table(db_path, table_name)
    columns, rows = _fetch_table(db_path, table_name)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = table_name[:31]  # Excel sheet-name limit

    # Header row — bold
    ws.append(columns)
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold

    # Data rows
    for row in rows:
        ws.append(list(row))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def export_to_json(db_path: str, table_name: str) -> bytes:
    """Export *table_name* from *db_path* as a UTF-8 JSON array of objects.

    Each row becomes a JSON object whose keys are column names.  Values
    that are not natively JSON-serialisable (Decimal, date, datetime, …)
    are converted to their string/float equivalents.

    Parameters
    ----------
    db_path:
        Absolute path to the workspace ``.db`` file.
    table_name:
        Name of the table to export.

    Returns
    -------
    bytes
        UTF-8-encoded JSON array.

    Raises
    ------
    ValueError
        If *db_path* does not exist or *table_name* is not present.
    """
    _assert_db_and_table(db_path, table_name)
    columns, rows = _fetch_table(db_path, table_name)

    data: list[dict[str, Any]] = [
        {col: _to_json_safe(val) for col, val in zip(columns, row)}
        for row in rows
    ]
    return json.dumps(data, default=str).encode("utf-8")
